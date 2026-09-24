#!/usr/bin/env python3
import argparse
import json
import re
import sys
import unicodedata
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

NAMESPACE = uuid.UUID('1b8b8d4e-b783-4ddd-8b25-d669fd44da90')
DIRECT_MEDIA = re.compile(r'\.(webp|gif|jpe?g|png|mp4)(?:\?.*)?$', re.I)


def uid(kind: str, key: str) -> str:
    return str(uuid.uuid5(NAMESPACE, f'{kind}:{key}'))


def text(value: Any) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def slugify(value: str) -> str:
    value = unicodedata.normalize('NFKD', value).encode('ascii', 'ignore').decode('ascii')
    value = re.sub(r'[^a-zA-Z0-9]+', '-', value).strip('-').lower()
    return value or 'item'


def normalize_space(value: str | None) -> str | None:
    if value is None:
        return None
    return re.sub(r'\s+', ' ', value).strip()


def parse_seconds(value: Any) -> int | None:
    s = normalize_space(text(value))
    if not s:
        return None
    low = s.lower()
    if 'no rest' in low:
        return 0
    m = re.search(r'(?:(\d+)\s*min)\s*(\d+)?', low)
    if m:
        mins = int(m.group(1))
        tail = int(m.group(2) or 0)
        return mins * 60 + tail
    m = re.search(r'(\d+)\s*sec', low)
    if m:
        return int(m.group(1))
    return None


def parse_prescription(value: Any) -> dict[str, Any]:
    raw = normalize_space(text(value))
    out: dict[str, Any] = {
        'raw': raw,
        'sets': None,
        'reps': None,
        'duration_sec': None,
        'laterality': None,
        'mode': None,
    }
    if not raw:
        return out
    low = raw.lower().replace(',', '.')
    if 'amrap' in low:
        out['mode'] = 'amrap'
        out['duration_sec'] = parse_seconds(low)
        return out
    m = re.match(r'^(\d+)\s*[xX]\s*(\d+)\s*sec(?:\s*/\s*([a-zA-Z]+))?', low)
    if m:
        out['sets'] = int(m.group(1))
        out['duration_sec'] = int(m.group(2))
        out['laterality'] = m.group(3)
        return out
    m = re.match(r'^(\d+)\s*[xX]\s*(\d+)(?:\s*/\s*([^\s]+))?', low)
    if m:
        out['sets'] = int(m.group(1))
        out['reps'] = int(m.group(2))
        out['laterality'] = m.group(3)
        return out
    m = re.match(r'^(\d+)\s*/\s*([^\s]+)', low)
    if m:
        out['sets'] = 1
        out['reps'] = int(m.group(1))
        out['laterality'] = m.group(2)
        return out
    if 'passage' in low:
        m = re.search(r'(\d+)', low)
        if m:
            out['sets'] = 1
            out['reps'] = int(m.group(1))
            return out
    duration = parse_seconds(low)
    if duration is not None:
        out['sets'] = 1
        out['duration_sec'] = duration
        return out
    if isinstance(value, (int, float)):
        out['sets'] = 1
        out['reps'] = int(value) if float(value).is_integer() else float(value)
        return out
    m = re.match(r'^(\d+)(?:\.0)?$', low)
    if m:
        out['sets'] = 1
        out['reps'] = int(m.group(1))
    return out


def parse_load(value: Any, sheet_default_unit: str | None) -> dict[str, Any]:
    raw = normalize_space(text(value))
    out: dict[str, Any] = {
        'raw': raw,
        'value': None,
        'unit': None,
        'quantity': None,
        'ambiguous': False,
        'cardio': None,
    }
    if not raw:
        return out
    low = raw.lower().replace(',', '.')
    speed = re.search(r'(\d+(?:\.\d+)?)\s*km\s*/?\s*h', low)
    incline = re.search(r'(?:pente\s*)?(\d+(?:\.\d+)?)\s*%', low)
    level = re.search(r'niveau\s*(\d+(?:\.\d+)?)', low)
    if speed or incline or level:
        out['cardio'] = {
            'speed_kmh': float(speed.group(1)) if speed else None,
            'incline_pct': float(incline.group(1)) if incline else None,
            'level': float(level.group(1)) if level else None,
        }
        return out
    if 'pas de charge' in low or 'barre seule' in low and 'kg' not in low:
        return out
    unit = None
    if 'lbs' in low or re.search(r'\blb\b', low):
        unit = 'lbs'
    elif 'kg' in low:
        unit = 'kg'
    nums = re.findall(r'\d+(?:\.\d+)?', low)
    if nums:
        out['value'] = float(nums[0])
        q = re.search(r'[xX]\s*(\d+)\b', raw)
        out['quantity'] = int(q.group(1)) if q else None
        if unit:
            out['unit'] = unit
        elif sheet_default_unit:
            out['unit'] = sheet_default_unit
        else:
            out['ambiguous'] = True
    return out


def media_record(exercise_id: str, url: str, source_sheet: str, source_cell: str) -> dict[str, Any]:
    ext = DIRECT_MEDIA.search(url)
    if ext:
        suffix = ext.group(1).lower()
        media_type = 'video' if suffix == 'mp4' else 'animated_image' if suffix == 'gif' else 'image'
        status = 'remote_direct'
    else:
        media_type = 'external_reference'
        status = 'reference_only'
    return {
        'id': uid('media', f'{exercise_id}:{url}'),
        'exercise_id': exercise_id,
        'media_type': media_type,
        'storage_path': None,
        'external_url': url,
        'original_source_url': url,
        'attribution': 'LiftManual source link preserved from the coach workbook',
        'alt_text': None,
        'sort_order': 0,
        'is_primary': media_type != 'external_reference',
        'status': status,
        'source_sheet': source_sheet,
        'source_cell': source_cell,
    }


def classify_section(title: str) -> str:
    low = title.lower()
    if 'warm' in low:
        return 'warm_up'
    if 'cardio' in low or 'zone 2' in low:
        return 'cardio'
    if 'finisher' in low:
        return 'finisher'
    if 'abs' in low or 'core' in low:
        return 'abs_core'
    if 'lower' in low:
        return 'lower_focus'
    if 'upper' in low:
        return 'upper_focus'
    if 'glute' in low:
        return 'glutes_focus'
    return 'workout'


def canonical_exercise(name: str) -> tuple[str, str | None]:
    low = normalize_space(name) or name
    test_ref = None
    if 'back squat test' in low.lower():
        return 'Barbell Back Squat', 'back-squat'
    if 'hip thrust test' in low.lower():
        return 'Hip Thrust', 'hip-thrust'
    return low, test_ref


def extract_hyperlink(cell: Any) -> tuple[str | None, str | None]:
    if not cell.hyperlink:
        return None, None
    return getattr(cell.hyperlink, 'target', None), getattr(cell.hyperlink, 'location', None)


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser()
    ap.add_argument('input')
    ap.add_argument('output')
    ap.add_argument('--report', default=None)
    args = ap.parse_args()
    input_path = Path(args.input)
    output_path = Path(args.output)
    report_path = Path(args.report) if args.report else output_path.with_name('import-report.json')

    wbf = load_workbook(input_path, data_only=False)
    wbv = load_workbook(input_path, data_only=True)

    dash = wbf['\U0001f3c5 Dashboard']
    dashv = wbv['\U0001f3c5 Dashboard']
    profile = {
        'display_name': text(dashv['B5'].value),
        'start_date': text(dashv['B6'].value),
        'level': text(dashv['B7'].value),
        'coach_calorie_context': text(dashv['B8'].value),
        'source_title': text(dashv['A1'].value),
    }

    phases = []
    for r in range(25, 28):
        phases.append({
            'id': uid('phase', str(r)),
            'name': text(dashv[f'A{r}'].value),
            'week_range': text(dashv[f'B{r}'].value),
            'label': text(dashv[f'D{r}'].value),
            'description': text(dashv[f'F{r}'].value),
            'sort_order': r - 24,
            'source_sheet': dash.title,
            'source_cell': f'A{r}:L{r}',
        })

    source_1rm_inputs = []
    for r in range(14, 17):
        source_1rm_inputs.append({
            'exercise': text(dashv[f'A{r}'].value),
            'input_kg': dashv[f'E{r}'].value,
            'percentages': {str(p): dashv.cell(r, c).value for p, c in [(50,7),(60,8),(70,9),(75,10),(80,11),(85,12)]},
            'formulas': {str(p): wbf[dash.title].cell(r, c).value for p, c in [(50,7),(60,8),(70,9),(75,10),(80,11),(85,12)]},
            'source_cell': f'A{r}:L{r}',
        })

    exercises: dict[str, dict[str, Any]] = {}
    media: dict[str, dict[str, Any]] = {}
    weeks = []
    anomalies: list[dict[str, Any]] = []

    for week_number in range(1, 5):
        sheet_name = f'Week {week_number}'
        ws = wbf[sheet_name]
        wsv = wbv[sheet_name]
        header = text(wsv['E3'].value) or ''
        sheet_default_unit = 'lbs' if 'LBS' in header.upper() else None
        week = {
            'id': uid('week', str(week_number)),
            'week_number': week_number,
            'name': sheet_name,
            'status': 'available',
            'source_sheet': sheet_name,
            'days': [],
        }
        current_day: dict[str, Any] | None = None
        current_section: dict[str, Any] | None = None
        for r in range(5, 71):
            aval = normalize_space(text(wsv.cell(r,1).value))
            bval = normalize_space(text(wsv.cell(r,2).value))
            if aval and re.match(r'^Jour\s+\d+$', aval, re.I):
                day_number = int(re.search(r'(\d+)', aval).group(1))
                current_day = {
                    'id': uid('day', f'{week_number}:{day_number}'),
                    'day_number': day_number,
                    'title': f'Jour {day_number}',
                    'description': None,
                    'estimated_duration_min': None,
                    'sort_order': day_number,
                    'source_sheet': sheet_name,
                    'source_cell': f'A{r}',
                    'sections': [],
                }
                week['days'].append(current_day)
                current_section = None
                continue
            if aval == 'POINTS DE VIGILANCE':
                current_day = None
                current_section = None
                continue
            if current_day is None:
                continue
            if aval and aval not in {'Ma partie a ne pas modifier'}:
                current_section = {
                    'id': uid('section', f'{week_number}:{current_day["day_number"]}:{r}:{aval}'),
                    'section_type': classify_section(aval),
                    'title': aval,
                    'instructions': None,
                    'sort_order': len(current_day['sections']) + 1,
                    'source_sheet': sheet_name,
                    'source_cell': f'A{r}',
                    'items': [],
                }
                current_day['sections'].append(current_section)
            if not bval:
                continue
            if current_section is None:
                current_section = {
                    'id': uid('section', f'{week_number}:{current_day["day_number"]}:misc'),
                    'section_type': 'workout',
                    'title': 'Workout',
                    'instructions': None,
                    'sort_order': len(current_day['sections']) + 1,
                    'source_sheet': sheet_name,
                    'source_cell': f'B{r}',
                    'items': [],
                }
                current_day['sections'].append(current_section)
            display_name, test_ref = canonical_exercise(bval)
            ex_slug = slugify(display_name)
            ex_id = uid('exercise', ex_slug)
            if ex_id not in exercises:
                exercises[ex_id] = {
                    'id': ex_id,
                    'slug': ex_slug,
                    'name': display_name,
                    'description': None,
                    'equipment': None,
                    'muscle_group': None,
                    'default_rest_seconds': None,
                    'aliases': [],
                }
            if bval != display_name and bval not in exercises[ex_id]['aliases']:
                exercises[ex_id]['aliases'].append(bval)
            source_url, source_location = extract_hyperlink(ws.cell(r,2))
            if source_url:
                rec = media_record(ex_id, source_url, sheet_name, f'B{r}')
                media[rec['id']] = rec
            cval = wsv.cell(r,3).value
            dval = wsv.cell(r,4).value
            eval_ = wsv.cell(r,5).value
            fval = normalize_space(text(wsv.cell(r,6).value))
            gval = normalize_space(text(wsv.cell(r,7).value))
            prescription = parse_prescription(cval)
            if current_section['section_type'] == 'cardio' and prescription['duration_sec'] is None:
                section_duration = parse_seconds(current_section['title'])
                if section_duration is not None:
                    prescription['duration_sec'] = section_duration
            load = parse_load(eval_, sheet_default_unit)
            cardio_from_target = parse_load(cval, None).get('cardio')
            if cardio_from_target and not load.get('cardio'):
                load['cardio'] = cardio_from_target
            item = {
                'id': uid('item', f'{sheet_name}:{r}'),
                'exercise_id': ex_id,
                'exercise_name': display_name,
                'item_kind': 'strength_test' if test_ref else ('cardio' if current_section['section_type'] == 'cardio' else 'exercise'),
                'strength_test_ref': test_ref,
                'prescribed_sets': prescription['sets'],
                'prescribed_reps': prescription['reps'],
                'prescribed_duration_sec': prescription['duration_sec'],
                'prescription_mode': prescription['mode'],
                'laterality': prescription['laterality'],
                'target_raw': prescription['raw'],
                'prescribed_weight': load['value'],
                'weight_unit': load['unit'],
                'weight_quantity': load['quantity'],
                'load_raw': load['raw'],
                'rest_seconds': parse_seconds(dval),
                'rest_raw': normalize_space(text(dval)),
                'intensity': None,
                'cardio': load['cardio'],
                'notes': fval,
                'athlete_comment_source': gval,
                'sort_order': len(current_section['items']) + 1,
                'source_sheet': sheet_name,
                'source_cell': f'B{r}:G{r}',
                'source_url': source_url,
                'source_location': source_location,
            }
            if load['ambiguous']:
                anomalies.append({
                    'code': 'AMBIGUOUS_UNIT',
                    'severity': 'medium',
                    'message': f'Load has a number but no explicit unit in {sheet_name} row {r}.',
                    'source_sheet': sheet_name,
                    'source_cell': f'E{r}',
                    'source_value': eval_,
                })
            current_section['items'].append(item)
        week['vigilance_notes'] = [
            normalize_space(text(wsv.cell(r,1).value)) for r in range(67, 71) if text(wsv.cell(r,1).value)
        ]
        weeks.append(week)

    tests = []
    test_ws = wbf['TESTS']
    test_wsv = wbv['TESTS']
    headers = []
    for r in range(1, 80):
        aval = normalize_space(text(test_wsv.cell(r,1).value))
        if aval in {'BACK SQUAT', 'HIP THRUST', 'DEADLIFT'}:
            headers.append((r, aval))
    for idx, (start_row, name) in enumerate(headers):
        end_row = headers[idx+1][0] - 1 if idx + 1 < len(headers) else 80
        slug = slugify(name)
        ex_name = 'Barbell Back Squat' if name == 'BACK SQUAT' else name.title()
        ex_id = uid('exercise', slugify(ex_name))
        exercises.setdefault(ex_id, {
            'id': ex_id,
            'slug': slugify(ex_name),
            'name': ex_name,
            'description': None,
            'equipment': None,
            'muscle_group': None,
            'default_rest_seconds': 120,
            'aliases': [],
        })
        source_url, _ = extract_hyperlink(test_ws.cell(start_row,1))
        if source_url:
            rec = media_record(ex_id, source_url, 'TESTS', f'A{start_row}')
            media[rec['id']] = rec
        sets = []
        formula = None
        formula_cached = None
        for r in range(start_row + 1, end_row + 1):
            aval = normalize_space(text(test_wsv.cell(r,1).value))
            if aval and aval.lower().startswith('serie'):
                raw_weight = normalize_space(text(test_wsv.cell(r,2).value))
                load = parse_load(raw_weight, 'kg')
                sets.append({
                    'set_number': int(re.search(r'(\d+)', aval).group(1)),
                    'weight_raw': raw_weight,
                    'weight': load['value'],
                    'weight_unit': load['unit'],
                    'repetitions': test_wsv.cell(r,3).value,
                    'rest_seconds': parse_seconds(test_wsv.cell(r,4).value),
                    'rest_raw': normalize_space(text(test_wsv.cell(r,4).value)),
                    'source_cell': f'A{r}:D{r}',
                })
            if aval and '1RM' in aval:
                formula = test_ws.cell(r,2).value
                formula_cached = test_wsv.cell(r,2).value
        tests.append({
            'id': uid('strength-test', slug),
            'slug': slug,
            'name': name.title(),
            'exercise_id': ex_id,
            'formula_name': 'Brzycki' if formula else None,
            'formula': formula,
            'source_estimated_1rm': formula_cached,
            'sets': sets,
            'source_sheet': 'TESTS',
            'source_cell': f'A{start_row}:D{end_row}',
        })

    # Known workbook contradictions are preserved and surfaced.
    anomalies.extend([
        {
            'code': 'PROGRAM_LENGTH_CONTRADICTION',
            'severity': 'high',
            'message': 'Dashboard title says 12 weeks, while phase ranges extend through week 13.',
            'source_sheet': dash.title,
            'source_cell': 'A1;B25:B27',
            'source_value': {'title': dashv['A1'].value, 'phase_ranges': [p['week_range'] for p in phases]},
        },
        {
            'code': 'ONLY_FOUR_WEEK_SHEETS',
            'severity': 'high',
            'message': 'Only Week 1 through Week 4 worksheets are present even though later phases are described.',
            'source_sheet': None,
            'source_cell': None,
            'source_value': [s for s in wbf.sheetnames if s.startswith('Week ')],
        },
        {
            'code': 'BACK_SQUAT_SET_2_TEXT_MATH',
            'severity': 'medium',
            'message': 'Back Squat set 2 says 30 kg but also says 20 kg bar plus two 5 kg plates on each side, which would total 40 kg.',
            'source_sheet': 'TESTS',
            'source_cell': 'B10',
            'source_value': test_wsv['B10'].value,
        },
        {
            'code': 'DEADLIFT_TEST_WEIGHTS_MISSING',
            'severity': 'medium',
            'message': 'Deadlift test includes reps and rest but no prescribed loads.',
            'source_sheet': 'TESTS',
            'source_cell': 'B36:B41',
            'source_value': None,
        },
        {
            'code': 'WEEK_2_LOADS_MISSING',
            'severity': 'medium',
            'message': 'Week 2 retains prescriptions and coaching notes but most load cells are blank.',
            'source_sheet': 'Week 2',
            'source_cell': 'E15:E60',
            'source_value': None,
        },
    ])

    result = {
        'schema_version': 1,
        'source': {
            'filename': input_path.name,
            'workbook_sheets': wbf.sheetnames,
            'importer': 'scripts/import_training_xlsx.py',
        },
        'athlete': {
            'id': uid('athlete', profile['display_name'] or 'kemi'),
            **profile,
            'preferred_units': 'kg',
            'locale': 'fr-CA',
        },
        'program': {
            'id': uid('program', 'kemi-training'),
            'name': 'KEMI Training',
            'description': 'Training program normalized from the coach workbook without changing prescribed content.',
            'start_date': profile['start_date'],
            'status': 'active',
            'source': 'xlsx',
            'phases': phases,
            'weeks': weeks,
            'vigilance_notes': weeks[0].get('vigilance_notes', []) if weeks else [],
        },
        'exercises': sorted(exercises.values(), key=lambda x: x['name'].lower()),
        'exercise_media': sorted(media.values(), key=lambda x: (x['exercise_id'], x['sort_order'], x['id'])),
        'strength_tests': tests,
        'source_1rm_inputs': source_1rm_inputs,
        'anomalies': anomalies,
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8', newline='\n')
    report = {
        'source': input_path.name,
        'sheet_count': len(wbf.sheetnames),
        'week_count': len(weeks),
        'workout_day_count': sum(len(w['days']) for w in weeks),
        'workout_item_count': sum(len(s['items']) for w in weeks for d in w['days'] for s in d['sections']),
        'exercise_count': len(exercises),
        'media_count': len(media),
        'direct_media_count': sum(1 for m in media.values() if m['media_type'] != 'external_reference'),
        'external_reference_count': sum(1 for m in media.values() if m['media_type'] == 'external_reference'),
        'strength_test_count': len(tests),
        'anomalies': anomalies,
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8', newline='\n')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
